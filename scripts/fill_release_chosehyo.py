#!/usr/bin/env python3
"""
リリース調整表 (リリース一覧_サマリ) への追記スクリプト。

Usage:
    python fill_release_chosehyo.py <template_path> <output_path> '<json>'

    output_path が存在しない場合は template_path をコピーして作成する。
    存在する場合はそのファイルに追記する。

JSON keys:
    kojin_dantai  : str  "個保" | "団保"
    release_type  : str  "開発" | "定例" | "保守" | "障害"
    anken         : str  案件名
    system        : str  アプリケーション・システム名
    tanto_shain   : str  リリース担当社員
    tanto_team    : str  リリース担当開発チーム
    kankyo        : str  環境 (例: UAT / PROD)
    yobi          : str  予備（空文字可）
    date          : str  日付 YYYY-MM-DD
    sakugyou_team : str  作業チーム
    sakugyou_time : str  作業時間 (例: 1h)
    start_time    : str  開始時刻 (例: 13:00)
    menus         : list 選択されたリリースメニュー名のリスト

Exit codes: 0=success, 1=error
"""

import json
import shutil
import sys
from datetime import datetime
from pathlib import Path
import openpyxl

MENU_COLUMNS = {
    "blobインサート":                 "O",
    "DBパッチ(SQL実行)":              "P",
    "DBパッチ(Oneshotジョブ実行)":    "Q",
    "ディレクトリ削除(Linux)":        "R",
    "ディレクトリ作成(Linux)":        "S",
    "A-AUTO実行(フォルダ作成バッチ)": "T",
    "JP1":                            "U",
    "サービス開閉局":                 "V",
    "サービス停止":                   "W",
    "OPX資材配置":                    "X",
    "ファイルアップロード(差し替え)":  "Y",
    "jenkinsモジュールリリース":      "Z",
    "CI/CDモジュール":                "AA",
    "その他":                         "AB",
}

DATA_START_ROW = 7


def find_next_empty_row(ws):
    for row_idx in range(DATA_START_ROW, ws.max_row + 2):
        if ws.cell(row=row_idx, column=2).value is None:
            return row_idx
    return ws.max_row + 1


def col_letter_to_index(letter):
    result = 0
    for ch in letter.upper():
        result = result * 26 + (ord(ch) - ord('A') + 1)
    return result


def main():
    if len(sys.argv) != 4:
        print("Usage: fill_release_chosehyo.py <template_path> <output_path> '<json>'", file=sys.stderr)
        sys.exit(1)

    template_path = sys.argv[1]
    output_path = sys.argv[2]
    data = json.loads(sys.argv[3])

    if not Path(output_path).exists():
        shutil.copy2(template_path, output_path)

    wb = openpyxl.load_workbook(output_path)
    ws = wb["リリース一覧_サマリ"]

    row = find_next_empty_row(ws)

    ws.cell(row=row, column=2).value = data["kojin_dantai"]
    ws.cell(row=row, column=3).value = data["release_type"]
    ws.cell(row=row, column=5).value = data["anken"]
    ws.cell(row=row, column=6).value = data["system"]
    ws.cell(row=row, column=7).value = data["tanto_shain"]
    ws.cell(row=row, column=8).value = data["tanto_team"]
    ws.cell(row=row, column=9).value = data["kankyo"]
    ws.cell(row=row, column=10).value = data.get("yobi") or None
    ws.cell(row=row, column=11).value = datetime.strptime(data["date"], "%Y-%m-%d")
    ws.cell(row=row, column=12).value = data["sakugyou_team"]
    ws.cell(row=row, column=13).value = data["sakugyou_time"]
    ws.cell(row=row, column=14).value = data["start_time"]

    for menu_name in data.get("menus", []):
        col_letter = MENU_COLUMNS.get(menu_name)
        if col_letter:
            ws.cell(row=row, column=col_letter_to_index(col_letter)).value = "●"
        else:
            print(f"Warning: unknown menu '{menu_name}'", file=sys.stderr)

    wb.save(output_path)
    print(json.dumps({"status": "ok", "row": row, "file": output_path}, ensure_ascii=False))


if __name__ == "__main__":
    main()
